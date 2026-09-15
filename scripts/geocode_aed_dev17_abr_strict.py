#!/usr/bin/env python3
"""Strictly geocode three address-only official AED datasets with ABR Geocoder.

This script only creates review evidence.  It never writes to Supabase and it
does not treat a geocoder result as publishable unless every strict condition
and the municipality review bounds pass.
"""
from __future__ import annotations

import hashlib
import json
import subprocess
import time
import urllib.parse
import urllib.request
from io import BytesIO
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from import_aed_open_data import (
    ADDRESS_FIELDS,
    NAME_FIELDS,
    PHONE_FIELDS,
    first_value,
    read_records,
)

RAW = Path("data/aed_dev14/raw")
OUT = Path("data/aed_dev17_abr_strict")

TARGETS = {
    "04207": {
        "prefecture": "宮城県",
        "municipality": "名取市",
        "path": RAW / "04207_32919.bin",
        "parser": "standard",
        "source_url": "https://miyagi.dataeye.jp/resources/32919",
        "source_license": "cc-by4_0",
        "source_updated_at": "2026-05-25",
        "bounds": (38.05, 38.30, 140.75, 141.10),
    },
    "12225": {
        "prefecture": "千葉県",
        "municipality": "君津市",
        "path": RAW / "12225_56087.bin",
        "parser": "kimitsu",
        "source_url": "https://opendata.pref.chiba.lg.jp/resources/56087",
        "source_license": None,
        "source_updated_at": "2024-07-29",
        "bounds": (35.10, 35.50, 139.70, 140.20),
    },
    "12228": {
        "prefecture": "千葉県",
        "municipality": "四街道市",
        "path": RAW / "12228_56525.bin",
        "parser": "yotsukaido",
        "source_url": "https://opendata.pref.chiba.lg.jp/resources/56525",
        "source_license": None,
        "source_updated_at": "2025-04-01",
        "bounds": (35.60, 35.75, 140.10, 140.30),
    },
}


def workbook_records(payload: bytes, sheet: str, header_row: int) -> list[dict[str, Any]]:
    ws = load_workbook(BytesIO(payload), read_only=True, data_only=True)[sheet]
    values = list(ws.iter_rows(values_only=True))
    headers = [str(value).strip() if value is not None else "" for value in values[header_row - 1]]
    records = []
    for row in values[header_row:]:
        record = {headers[i]: value for i, value in enumerate(row) if headers[i]}
        if any(value not in (None, "") for value in record.values()):
            records.append(record)
    return records


def load_records(target: dict[str, Any]) -> list[dict[str, Any]]:
    payload = target["path"].read_bytes()
    if target["parser"] == "standard":
        return read_records(payload)
    if target["parser"] == "kimitsu":
        return workbook_records(payload, "AED設置箇所一覧", 2)
    if target["parser"] == "yotsukaido":
        rows = workbook_records(payload, "1113", 3)
        for row in rows:
            row["名称"] = row.get("機関・施設名")
            row["所在地"] = row.get("住所")
        return rows
    raise ValueError(f"unknown parser: {target['parser']}")


def geocode(address: str) -> dict[str, Any] | None:
    url = "http://127.0.0.1:3000/geocode?" + urllib.parse.urlencode(
        {"address": address, "target": "residential", "format": "json"}
    )
    try:
        with urllib.request.urlopen(url, timeout=20) as response:
            payload = json.loads(response.read().decode("utf-8"))
        return payload[0] if payload else None
    except Exception:
        return None


def wait_server() -> None:
    for _ in range(90):
        try:
            with urllib.request.urlopen("http://127.0.0.1:3000/health", timeout=2) as response:
                if response.status == 200:
                    return
        except Exception:
            time.sleep(1)
    raise RuntimeError("ABR geocoder server did not become healthy")


def process(code: str, target: dict[str, Any]) -> None:
    prefecture = target["prefecture"]
    municipality = target["municipality"]
    rows = load_records(target)
    dbdir = Path("/tmp") / f"abr-dev17-{code}"
    subprocess.run(["abrg", "download", "-c", code, "-d", str(dbdir), "--silent"], check=True)
    subprocess.run(["abrg", "serve", "start", "-d", str(dbdir)], check=True)
    wait_server()
    accepted: list[dict[str, Any]] = []
    held: list[dict[str, Any]] = []
    try:
        for row_number, raw in enumerate(rows, 2):
            name = str(first_value(raw, NAME_FIELDS) or "").strip()
            address = str(first_value(raw, ADDRESS_FIELDS) or "").strip()
            phone = str(first_value(raw, PHONE_FIELDS) or "").strip() or None
            if address and not address.startswith(prefecture):
                address = prefecture + (address if address.startswith(municipality) else municipality + address)
            base = {
                "row": row_number,
                "name": name,
                "address": address,
                "phone": phone,
                "source_url": target["source_url"],
                "source_license": target["source_license"],
                "source_updated_at": target["source_updated_at"],
            }
            if not name or not address:
                held.append({**base, "reason": "missing_name_or_address"})
                continue
            response = geocode(address)
            result = (response or {}).get("result") or {}
            lat = result.get("lat")
            lon = result.get("lon")
            south, north, west, east = target["bounds"]
            in_bounds = (
                lat is not None
                and lon is not None
                and south <= float(lat) <= north
                and west <= float(lon) <= east
            )
            strict = (
                result.get("match_level") == "residential_detail"
                and result.get("coordinate_level") == "residential_detail"
                and float(result.get("score") or 0) >= 0.90
                and str(result.get("lg_code") or "") == code
                and not (result.get("others") or [])
                and in_bounds
            )
            detail = {
                "score": result.get("score"),
                "match_level": result.get("match_level"),
                "coordinate_level": result.get("coordinate_level"),
                "lg_code": result.get("lg_code"),
                "others": result.get("others"),
                "normalized_address": result.get("output"),
            }
            if strict:
                accepted.append(
                    {
                        **base,
                        **detail,
                        "latitude": float(lat),
                        "longitude": float(lon),
                    }
                )
            else:
                held.append({**base, **detail, "reason": "strict_residential_match_failed"})
    finally:
        subprocess.run(["abrg", "serve", "stop"], check=False)

    OUT.mkdir(parents=True, exist_ok=True)
    stem = f"{code}_{municipality}"
    (OUT / f"{stem}_accepted.json").write_text(
        json.dumps(accepted, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    (OUT / f"{stem}_held.json").write_text(
        json.dumps(held, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    report = {
        "code": code,
        "prefecture": prefecture,
        "municipality": municipality,
        "source_rows": len(rows),
        "accepted": len(accepted),
        "held": len(held),
        "license_verified": target["source_license"] is not None,
        "source_sha256": hashlib.sha256(target["path"].read_bytes()).hexdigest(),
        "policy": "match_level=residential_detail; coordinate_level=residential_detail; score>=0.90; lg_code exact; others empty; municipality bounds",
    }
    (OUT / f"{stem}_report.json").write_text(
        json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8"
    )
    print(code, municipality, "accepted", len(accepted), "held", len(held), flush=True)


def main() -> None:
    for code, target in TARGETS.items():
        process(code, target)


if __name__ == "__main__":
    main()
